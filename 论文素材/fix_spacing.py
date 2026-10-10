from pathlib import Path
p=Path('生成论文.py');s=p.read_text(encoding='utf-8-sig')
s=s.replace('TA_CENTER, TA_JUSTIFY','TA_CENTER, TA_JUSTIFY, TA_LEFT')
s=s.replace("font='C:/Windows/Fonts/NotoSansSC-VF.ttf'","font='C:/Windows/Fonts/simsun.ttc'")
s=s.replace("fontName='NotoSC',**kw","fontName='NotoSC',wordWrap='CJK',**kw")
s=s.replace('alignment=TA_JUSTIFY','alignment=TA_LEFT')
s=s.replace("spaceAfter=8,textColor=colors.black","spaceAfter=8,keepWithNext=True,textColor=colors.black")
s=s.replace("spaceAfter=5,textColor=colors.black","spaceAfter=5,keepWithNext=True,textColor=colors.black")
s=s.replace('fontSize=10.2,leading=18','fontSize=10.5,leading=18')
s=s[:s.index('def footer')]
s += '''
# Both deliverables use the original semantic paragraphs, never PDF-extracted text.
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
D=Document(); sec=D.sections[0]
sec.page_width=Inches(210/25.4);sec.page_height=Inches(297/25.4)
sec.left_margin=sec.right_margin=Inches(18/25.4)
sec.top_margin=sec.bottom_margin=Inches(16/25.4)
for name,sz,east in [('Normal',10.5,'宋体'),('Title',20,'黑体'),('Heading 1',15,'黑体'),('Heading 2',12.5,'黑体'),('Caption',9,'宋体')]:
 st=D.styles[name]; st.font.name='Times New Roman';st.font.size=Pt(sz);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
 st.paragraph_format.page_break_before=False
 st.paragraph_format.first_line_indent=Pt(0)
 st.paragraph_format.space_before=Pt(6 if 'Heading' in name else 0)
 st.paragraph_format.space_after=Pt(5)
 st.paragraph_format.line_spacing=Pt(18)
 st.paragraph_format.keep_with_next=('Heading' in name)
 D.styles['Normal'].paragraph_format.first_line_indent=Pt(21)
for item in story:
 if isinstance(item,Paragraph):
  sn=item.style.name;txt=item.getPlainText()
  style={'Title':'Title','H1':'Heading 1','H2':'Heading 2','Caption':'Caption'}.get(sn,'Normal')
  p=D.add_paragraph(txt,style)
  p.alignment=WD_ALIGN_PARAGRAPH.LEFT
  if sn in ['Title','Sub','Caption']:
   p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0)
  if sn=='Title':p.paragraph_format.line_spacing=Pt(28)
  if sn=='Body0':p.paragraph_format.first_line_indent=Pt(0)
 elif isinstance(item,Table):
  vals=item._cellvalues;t=D.add_table(rows=len(vals),cols=len(vals[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER
  for i,row in enumerate(vals):
   for j,val in enumerate(row):
    c=t.cell(i,j);c.text=val.getPlainText() if isinstance(val,Paragraph) else str(val)
    for pp in c.paragraphs:
     pp.paragraph_format.first_line_indent=Pt(0);pp.paragraph_format.space_after=Pt(3);pp.paragraph_format.line_spacing=Pt(13)
     for rr in pp.runs:rr.font.size=Pt(9)
  pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
  for side in ['top','bottom']:
   el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'8');b.append(el)
  pr.append(b)
 elif isinstance(item,Image):
  p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.keep_with_next=True
  p.add_run().add_picture(item.filename,width=Pt(item.drawWidth),height=Pt(item.drawHeight))
f=sec.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)
D.save('占位学号-占位姓名-机器学习算法课程论文.docx')
# PDF table cells must wrap instead of overflow.
for item in story:
 if isinstance(item,Table):
  for i,row in enumerate(item._cellvalues):
   for j,val in enumerate(row):
    if not isinstance(val,Paragraph):item._cellvalues[i][j]=Paragraph(str(val),ParagraphStyle('Cell',fontName='NotoSC',fontSize=8,leading=12,wordWrap='CJK'))
# Keep only each figure and caption together, never whole sections.
flow=[];i=0
while i<len(story):
 if isinstance(story[i],Image) and i+1<len(story) and isinstance(story[i+1],Paragraph):
  flow.append(KeepTogether(story[i:i+2]));i+=2
 else:flow.append(story[i]);i+=1
def footer(canvas,doc):
 canvas.setFont('NotoSC',8);canvas.drawCentredString(A4[0]/2,9*mm,str(doc.page))
doc=SimpleDocTemplate('占位学号-占位姓名-机器学习算法课程论文.pdf',pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,topMargin=16*mm,bottomMargin=16*mm)
doc.build(flow,onFirstPage=footer,onLaterPages=footer)
print('Original-source DOCX and CJK-aware PDF regenerated.')
'''
p.write_text(s,encoding='utf-8')
