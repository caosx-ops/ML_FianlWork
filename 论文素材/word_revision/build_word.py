from pathlib import Path
import json, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
WORK=ROOT/'论文素材/word_revision'
DATA=json.loads((WORK/'supplementary.json').read_text(encoding='utf-8'))
RAW=json.loads((ROOT/'實驗結果/raw_predictions.json').read_text(encoding='utf-8'))
plt.rcParams.update({'font.family':'sans-serif','font.sans-serif':['SimSun'],'axes.unicode_minus':False,'font.size':10})
rows=DATA['controlled_ablation']
fig,ax=plt.subplots(figsize=(6.3,2.4))
ax.errorbar(range(5),[r['cv_mean']*100 for r in rows],yerr=[r['cv_std']*100 for r in rows],color='black',marker='o',capsize=4,linewidth=1)
ax.set_xticks(range(5),['无PCA','80%','90%','95%','99%']);ax.set_xlabel('方差保留阈值');ax.set_ylabel('五折准确率（%）');ax.grid(axis='y',alpha=.25);fig.tight_layout()
fig.savefig(WORK/'ablation.png',dpi=230);plt.close(fig)
cm=np.array(RAW['results'][0]['confusion_matrix']);fig,ax=plt.subplots(figsize=(5.2,3.2));ax.imshow(cm,cmap='Greys')
for i in range(10):
 for j in range(10):
  ax.text(j,i,str(cm[i,j]),ha='center',va='center',fontsize=8,color='white' if cm[i,j]>25 else 'black')
ax.set_xticks(range(10));ax.set_yticks(range(10));ax.set_xlabel('预测标签');ax.set_ylabel('真实标签');fig.tight_layout();fig.savefig(WORK/'confusion.png',dpi=230);plt.close(fig)

d=Document();s=d.sections[0];s.page_width=Mm(210);s.page_height=Mm(297);s.top_margin=s.bottom_margin=Mm(23);s.left_margin=s.right_margin=Mm(25);s.footer_distance=Mm(12)
for name,size,east in [('Normal',12,'宋体'),('Title',18,'黑体'),('Heading 1',14,'黑体'),('Heading 2',12,'黑体'),('Caption',10.5,'宋体')]:
 st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0)
 st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),east)
 pf=st.paragraph_format;pf.space_before=Pt(8 if 'Heading' in name else 0);pf.space_after=Pt(4 if 'Heading' in name else 0);pf.line_spacing=Pt(20);pf.first_line_indent=Pt(0);pf.page_break_before=False;pf.keep_with_next='Heading' in name;pf.widow_control=True
 d.styles['Normal'].paragraph_format.first_line_indent=Pt(24)
 d.styles['Normal'].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
# Avoid inherited display grids and automatic script spacing; no tabs/manual spaces for alignment.
grid=s._sectPr.find(qn('w:docGrid'))
if grid is not None:s._sectPr.remove(grid)
settings=d.settings.element
for name in ['doNotAutoHyphenate']:
 el=OxmlElement('w:'+name);settings.append(el)

def plain(text,style=None,indent=True):
 p=d.add_paragraph(text,style);p.alignment=WD_ALIGN_PARAGRAPH.LEFT
 if not indent:p.paragraph_format.first_line_indent=Pt(0)
 pp=p._p.get_or_add_pPr()
 for tag in ['autoSpaceDE','autoSpaceDN','snapToGrid','overflowPunct']:
  e=OxmlElement('w:'+tag);e.set(qn('w:val'),'0');pp.append(e)
 return p
def caption(text):
 p=plain(text,'Caption',False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(15);p.paragraph_format.space_after=Pt(4);return p
def border(cell,where,size):
 tc=cell._tc.get_or_add_tcPr();b=tc.find(qn('w:tcBorders'))
 if b is None:b=OxmlElement('w:tcBorders');tc.append(b)
 el=OxmlElement('w:'+where);el.set(qn('w:val'),'single');el.set(qn('w:sz'),str(size));b.append(el)
def table(title,values,widths):
 cap=caption(title);cap.paragraph_format.keep_with_next=True
 t=d.add_table(rows=len(values),cols=len(values[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Mm(w)
 for i,row in enumerate(values):
  trpr=t.rows[i]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
  if i==0:trpr.append(OxmlElement('w:tblHeader'))
  for j,value in enumerate(row):
   c=t.cell(i,j);c.width=Mm(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;c.text=str(value)
   p=c.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Pt(0);p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_before=Pt(2);p.paragraph_format.space_after=Pt(2)
   p.paragraph_format.keep_with_next=(i<len(values)-1)
   for r in p.runs:r.font.size=Pt(10.5);r.bold=i==0
   if i==0:border(c,'top',10);border(c,'bottom',5)
   if i==len(values)-1:border(c,'bottom',10)
def figure(file,title,width):
 p=plain('',indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.line_spacing=1.0;p.add_run().add_picture(str(file),width=Mm(width));caption(title)

for line in (WORK/'正文.md').read_text(encoding='utf-8').splitlines():
 if not line.strip():continue
 if line.startswith('# '):
  p=plain(line[2:],'Title',False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(25);p.paragraph_format.space_after=Pt(8)
 elif line.startswith('## '):plain(line[3:],'Heading 1',False)
 elif line.startswith('### '):plain(line[4:],'Heading 2',False)
 elif line.startswith('@author '):
  p=plain(line[8:],indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(8)
 elif line.startswith('@abstract '):plain(line[10:],'Heading 2',False)
 elif line.startswith('@keywords '):
  p=plain(line[10:],indent=False);p.paragraph_format.space_after=Pt(4)
 elif line.startswith('@equation '):
  p=plain('',indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.line_spacing=Pt(22)
  math=OxmlElement('m:oMath');run=OxmlElement('m:r');txt=OxmlElement('m:t');txt.text=line[10:];run.append(txt);math.append(run);p._p.append(math)
 elif line=='@flow':
  p=plain('分层划分 → 训练折标准化 → 可选PCA → RBF-SVM → 折内验证',indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  p=plain('选定参数 → 全训练集重拟合 → 固定测试集评估',indent=False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
  caption('图1  无测试信息泄漏的训练与评估流程')
 elif line=='@table_data':
  table('表1  数据与评估协议',[['项目','设置'],['图像与特征','8×8灰度图像，64维，像素值0至16'],['数据规模','1797个样本，10个类别'],['训练集 / 测试集','1347 / 450，按类别分层'],['训练集内验证','五折分层，打乱，随机种子42'],['数据预处理','折内标准化，可选PCA，不做图像增强']],[43,117])
 elif line=='@table_main':
  table('表2  固定测试集上的主比较结果',[['模型','维数','准确率/%','宏P/%','宏R/%','宏F1/%'],['M1 标准化SVM',64,'98.22','98.27','98.20','98.21'],['M2 PCA调参SVM',54,'98.22','98.27','98.20','98.21'],['M3 PCA默认SVM',40,'98.22','98.27','98.20','98.21']],[46,16,27,23,23,25])
 elif line=='@table_ablation':
  vals=[['PCA阈值','全训练集维数','维度减少/%','五折准确率/%']]
  for r in rows:vals.append(['无PCA' if r['ratio'] is None else f"{r['ratio']*100:.0f}%",r['dimensions'],f"{(64-r['dimensions'])/64*100:.2f}",f"{r['cv_mean']*100:.2f} ± {r['cv_std']*100:.2f}"])
  table('表3  固定C与gamma的PCA消融结果',vals,[32,38,38,52])
 elif line=='@figure_ablation':figure(WORK/'ablation.png','图2  固定参数下的交叉验证结果与折间标准差',140)
 elif line=='@figure_confusion':figure(WORK/'confusion.png','图3  基线在450个测试样本上的混淆矩阵',126)
 elif line.startswith('@ref '):
  p=plain(line[5:],indent=False);p.paragraph_format.line_spacing=Pt(16);p.paragraph_format.space_after=Pt(4)
  for r in p.runs:r.font.size=Pt(10.5)
 else:plain(line)
f=s.footer.paragraphs[0];f.alignment=WD_ALIGN_PARAGRAPH.CENTER;f.paragraph_format.first_line_indent=Pt(0)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');f._p.append(fld)
d.core_properties.title='PCA降维对手写数字识别性能的影响研究';d.core_properties.author=''
dest=ROOT/'占位学号-占位姓名-机器学习算法课程论文-规范修订版.docx';d.save(dest);print(dest)
